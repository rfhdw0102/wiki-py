import re
import shutil
import tempfile
from dataclasses import dataclass
from importlib import import_module
from pathlib import Path
from typing import Any

from wiki_agent.models import SourceKind


class ComplexDocumentError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedSection:
    ordinal: int
    heading: str | None
    text: str
    locator: dict[str, object]


@dataclass(frozen=True)
class ParsedDocument:
    title: str
    markdown: str
    sections: tuple[ParsedSection, ...]
    parser: str


class DocumentParser:
    def parse(self, path: Path, kind: SourceKind) -> ParsedDocument:
        if kind in {SourceKind.MARKDOWN, SourceKind.TEXT}:
            return self._parse_text(path, kind)
        try:
            if kind in {SourceKind.HTML, SourceKind.WEBSITE}:
                return self._parse_html(path)
            if kind is SourceKind.DOCX:
                return self._parse_docx(path)
            if kind is SourceKind.PDF:
                return self._parse_pdf(path)
        except ComplexDocumentError:
            return self._parse_with_docling(path, kind)
        raise ValueError(f"unsupported source kind {kind.value}")

    @staticmethod
    def _parse_text(path: Path, kind: SourceKind) -> ParsedDocument:
        content = path.read_text(encoding="utf-8-sig")
        markdown = content if kind is SourceKind.MARKDOWN else f"```\n{content}\n```"
        return ParsedDocument(
            title=path.name,
            markdown=markdown,
            sections=_split_sections(content),
            parser="text",
        )

    @staticmethod
    def _parse_html(path: Path) -> ParsedDocument:
        try:
            from bs4 import BeautifulSoup
        except ImportError as exc:
            raise RuntimeError(
                "Install the document-parsing extra to parse HTML sources"
            ) from exc
        content = path.read_text(encoding="utf-8-sig")
        soup = BeautifulSoup(content, "html.parser")
        if soup.find(["table", "math", "svg"]):
            raise ComplexDocumentError("complex HTML requires Docling")
        for element in soup(["script", "style", "noscript"]):
            element.decompose()
        markdown_lines: list[str] = []
        for element in soup.select("h1,h2,h3,h4,h5,h6,p,li,pre"):
            text = element.get_text(" ", strip=True)
            if not text:
                continue
            if element.name and element.name.startswith("h"):
                markdown_lines.append(f"{'#' * int(element.name[1])} {text}")
            elif element.name == "li":
                markdown_lines.append(f"- {text}")
            elif element.name == "pre":
                markdown_lines.append(f"```\n{text}\n```")
            else:
                markdown_lines.append(text)
        markdown = "\n\n".join(markdown_lines)
        return ParsedDocument(
            title=_html_title(soup) or path.name,
            markdown=markdown,
            sections=_split_sections(markdown),
            parser="beautifulsoup",
        )

    @staticmethod
    def _parse_docx(path: Path) -> ParsedDocument:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError(
                "Install the document-parsing extra to parse DOCX sources"
            ) from exc
        document = Document(str(path))
        if document.tables:
            raise ComplexDocumentError("DOCX tables require Docling")
        markdown_lines: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name.lower() if paragraph.style is not None else ""
            heading_match = re.fullmatch(r"heading\s+([1-6])", style)
            if heading_match:
                markdown_lines.append(f"{'#' * int(heading_match.group(1))} {text}")
            else:
                markdown_lines.append(text)
        markdown = "\n\n".join(markdown_lines)
        title = document.core_properties.title or path.name
        return ParsedDocument(
            title=title,
            markdown=markdown,
            sections=_split_sections(markdown),
            parser="python-docx",
        )

    @staticmethod
    def _parse_pdf(path: Path) -> ParsedDocument:
        try:
            pymupdf = import_module("pymupdf")
        except ImportError as exc:
            raise RuntimeError(
                "Install the document-parsing extra to parse PDF sources"
            ) from exc
        document = pymupdf.open(path)
        page_markdown: list[str] = []
        sections: list[ParsedSection] = []
        try:
            for page_index, page in enumerate(document):
                finder = getattr(page, "find_tables", None)
                if finder is not None and finder().tables:
                    raise ComplexDocumentError("PDF tables require Docling")
                text = page.get_text("text").strip()
                if not text:
                    continue
                page_markdown.append(f"## Page {page_index + 1}\n\n{text}")
                for section in _split_sections(text):
                    sections.append(
                        ParsedSection(
                            ordinal=len(sections),
                            heading=section.heading or f"Page {page_index + 1}",
                            text=section.text,
                            locator={
                                **section.locator,
                                "page": page_index + 1,
                            },
                        )
                    )
        finally:
            document.close()
        markdown = "\n\n".join(page_markdown)
        if len(markdown.strip()) < 40:
            raise ComplexDocumentError("scanned or text-sparse PDF requires Docling OCR")
        return ParsedDocument(
            title=path.name,
            markdown=markdown,
            sections=tuple(sections),
            parser="pymupdf",
        )

    @staticmethod
    def _parse_with_docling(path: Path, kind: SourceKind) -> ParsedDocument:
        try:
            from docling.document_converter import DocumentConverter
        except ImportError as exc:
            raise RuntimeError(
                "Complex documents require the document-parsing extra with Docling"
            ) from exc
        suffix = {
            SourceKind.PDF: ".pdf",
            SourceKind.DOCX: ".docx",
            SourceKind.HTML: ".html",
            SourceKind.WEBSITE: ".html",
        }.get(kind)
        if suffix is None:
            raise ValueError(f"Docling parser does not support source kind {kind.value}")
        with tempfile.TemporaryDirectory(prefix="wiki-parse-") as temp_dir:
            named_source = Path(temp_dir) / f"source{suffix}"
            shutil.copyfile(path, named_source)
            result = DocumentConverter().convert(named_source)
            markdown = result.document.export_to_markdown()
        return ParsedDocument(
            title=path.name,
            markdown=markdown,
            sections=_split_sections(markdown),
            parser="docling",
        )


def _html_title(soup: Any) -> str | None:
    title = soup.find("title")
    if title is None:
        return None
    value = title.get_text(" ", strip=True)
    return value or None


def _split_sections(content: str, max_characters: int = 4_000) -> tuple[ParsedSection, ...]:
    heading_pattern = re.compile(r"(?m)^(#{1,6})\s+(.+?)\s*$")
    headings = list(heading_pattern.finditer(content))
    ranges: list[tuple[str | None, int, int]] = []
    if not headings:
        ranges.append((None, 0, len(content)))
    else:
        if headings[0].start() > 0 and content[: headings[0].start()].strip():
            ranges.append((None, 0, headings[0].start()))
        for index, match in enumerate(headings):
            end = headings[index + 1].start() if index + 1 < len(headings) else len(content)
            ranges.append((match.group(2).strip(), match.start(), end))

    sections: list[ParsedSection] = []
    for heading, start, end in ranges:
        segment = content[start:end].strip()
        if not segment:
            continue
        offset = 0
        while offset < len(segment):
            chunk_end = min(offset + max_characters, len(segment))
            if chunk_end < len(segment):
                boundary = segment.rfind("\n\n", offset, chunk_end)
                if boundary > offset:
                    chunk_end = boundary
            text = segment[offset:chunk_end].strip()
            if text:
                sections.append(
                    ParsedSection(
                        ordinal=len(sections),
                        heading=heading,
                        text=text,
                        locator={
                            "character_start": start + offset,
                            "character_end": start + chunk_end,
                        },
                    )
                )
            offset = max(chunk_end, offset + 1)
    if not sections:
        raise ValueError("parsed document contains no text")
    return tuple(sections)
