from pathlib import Path

from docx import Document
from wiki_agent.models import SourceKind
from wiki_agent.sources.parser import DocumentParser


def test_markdown_parser_preserves_heading_sections(tmp_path: Path) -> None:
    source = tmp_path / "manual.md"
    source.write_text(
        "# Compressor\n\nRated speed is 3000 rpm.\n\n## Shutdown\n\nClose the valve.",
        encoding="utf-8",
    )
    parsed = DocumentParser().parse(source, SourceKind.MARKDOWN)
    assert [section.heading for section in parsed.sections] == ["Compressor", "Shutdown"]
    assert parsed.sections[0].locator["character_start"] == 0
    assert "Close the valve" in parsed.sections[1].text


def test_text_parser_rejects_empty_document(tmp_path: Path) -> None:
    source = tmp_path / "empty.txt"
    source.write_text("", encoding="utf-8")
    try:
        DocumentParser().parse(source, SourceKind.TEXT)
    except ValueError as exc:
        assert "contains no text" in str(exc)
    else:
        raise AssertionError("empty documents must be rejected")


def test_html_parser_uses_lightweight_parser_for_simple_content(tmp_path: Path) -> None:
    source = tmp_path / "manual.html"
    source.write_text(
        "<html><head><title>Manual</title></head><body>"
        "<h1>Startup</h1><p>Open the isolation valve.</p></body></html>",
        encoding="utf-8",
    )

    parsed = DocumentParser().parse(source, SourceKind.HTML)

    assert parsed.parser == "beautifulsoup"
    assert parsed.title == "Manual"
    assert parsed.sections[0].heading == "Startup"


def test_docx_parser_uses_lightweight_parser_without_tables(tmp_path: Path) -> None:
    source = tmp_path / "manual.docx"
    document = Document()
    document.core_properties.title = "Operations Manual"
    document.add_heading("Shutdown", level=1)
    document.add_paragraph("Close the isolation valve.")
    document.save(source)

    parsed = DocumentParser().parse(source, SourceKind.DOCX)

    assert parsed.parser == "python-docx"
    assert parsed.title == "Operations Manual"
    assert parsed.sections[0].heading == "Shutdown"
